import os, json, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_bank3 import RECORDS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

HEADER_FILL = PatternFill(start_color="1d4ed8", end_color="1d4ed8", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GOOD_FILL = PatternFill(start_color="c6f6d5", end_color="c6f6d5", fill_type="solid")
BAD_FILL = PatternFill(start_color="fed7d7", end_color="fed7d7", fill_type="solid")
MID_FILL = PatternFill(start_color="fefcbf", end_color="fefcbf", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="dc2626", end_color="dc2626", fill_type="solid")
THIN = Side(style="thin", color="cccccc")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("record", "Registro"), ("label", "Descripción"), ("prominence_factor", "prominence_factor"),
    ("n_gt_beats", "Latidos GT"), ("hr_ground_truth", "HR GT"),
    ("sens_amplitud_r2", "Sens. AMPLITUD (r2)"), ("ppv_amplitud_r2", "VPP AMPLITUD (r2)"),
    ("sens_pendiente_r3", "Sens. PENDIENTE (r3)"), ("ppv_pendiente_r3", "VPP PENDIENTE (r3)"),
    ("hr_error_r3", "Err.HR (NO usar solo)"), ("n_competitions", "Competencias"),
    ("confidence", "Confianza"), ("gate_decision", "Decisión compuerta"),
    ("is_paced_flag", "Bandera marcapasos (NO CONFIABLE)"), ("n_spikes", "n_espigas"),
]


def quality_fill(value):
    if value is None:
        return None
    if value >= 0.95:
        return GOOD_FILL
    if value >= 0.8:
        return MID_FILL
    return BAD_FILL


def build():
    rows = json.load(open(os.path.join(OUT_DIR, 'raw_results3.json')))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "3a ronda"

    ws.merge_cells('A1:O1')
    ws['A1'] = "Banco de validación 3ª ronda — Desempate por pendiente (Parte A) + bandera de marcapasos (Parte B, NO CONFIABLE)"
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells('A2:O2')
    ws['A2'] = "⚠ La columna 'Bandera marcapasos' es un resultado NEGATIVO documentado -- NO está cableada a la compuerta. Ver hoja de decisión."
    ws['A2'].font = Font(italic=True, size=9, color="c0392b")

    header_row = 4
    for i, (key, label) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BORDER

    r = header_row + 1
    for row in rows:
        for i, (key, label) in enumerate(COLUMNS, start=1):
            val = row[key]
            c = ws.cell(row=r, column=i, value=val)
            c.border = BORDER
            if key in ("sens_amplitud_r2", "ppv_amplitud_r2", "sens_pendiente_r3", "ppv_pendiente_r3"):
                fill = quality_fill(val)
                if fill:
                    c.fill = fill
            if key == "gate_decision":
                c.fill = GOOD_FILL if val == 'ACEPTA' else BAD_FILL
            if key == "is_paced_flag" and row['record_id'] != '217' and val is True:
                c.fill = CRITICAL_FILL
                c.font = Font(color="FFFFFF", bold=True)
            if key == "is_paced_flag" and row['record_id'] == '118' and val is True:
                c.fill = CRITICAL_FILL
                c.font = Font(color="FFFFFF", bold=True)
        r += 1

    widths = [10, 30, 14, 10, 9, 15, 14, 15, 14, 16, 12, 10, 14, 24, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # Hoja 2: resumen 118 antes/despues (recuperacion)
    ws2 = wb.create_sheet("118 -- Recuperación BRD")
    ws2['A1'] = "Registro 118 (bloqueo de rama derecha) -- desempate amplitud vs pendiente"
    ws2['A1'].font = Font(bold=True, size=13)
    headers2 = ["prominence_factor", "Sens. AMPLITUD (r2)", "VPP AMPLITUD (r2)", "Sens. PENDIENTE (r3)", "VPP PENDIENTE (r3)"]
    for i, h in enumerate(headers2, start=1):
        c = ws2.cell(row=3, column=i, value=h)
        c.fill = HEADER_FILL; c.font = HEADER_FONT; c.border = BORDER
        c.alignment = Alignment(wrap_text=True)
    r2 = 4
    for row in rows:
        if row['record_id'] == '118':
            vals = [row['prominence_factor'], row['sens_amplitud_r2'], row['ppv_amplitud_r2'], row['sens_pendiente_r3'], row['ppv_pendiente_r3']]
            for i, v in enumerate(vals, start=1):
                c = ws2.cell(row=r2, column=i, value=v)
                c.border = BORDER
                if i in (2,3):
                    fill = quality_fill(v)
                    if fill: c.fill = fill
                if i in (4,5):
                    fill = quality_fill(v)
                    if fill: c.fill = fill
            r2 += 1
    for i, w in enumerate([16,18,18,18,18], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # Hoja 3: evidencia del barrido de la bandera de marcapasos (Parte B, negativo)
    ws3 = wb.create_sheet("Evidencia Parte B (negativa)")
    ws3['A1'] = "Barrido de slope_z_threshold para la bandera de marcapasos -- ningún valor separa 217 de 118/ruido"
    ws3['A1'].font = Font(bold=True, size=12, color="c0392b")
    sweep_data = [
        ("z", "100", "103", "117", "108", "228", "207", "233", "118 (NUNCA debe activarse)", "210", "217 (debe activarse)"),
        (12, 0,0,0,54,2,12,6,11,2,10),
        (15, 0,0,0,33,11,2,5,32,1,1),
        (20, 0,0,0,13,6,3,0,0,0,0),
        (25, 0,0,0,3,0,3,5,0,0,0),
        (30, 0,0,0,3,0,5,10,0,7,3),
    ]
    for ri, row_data in enumerate(sweep_data, start=3):
        for ci, val in enumerate(row_data, start=1):
            c = ws3.cell(row=ri, column=ci, value=val)
            c.border = BORDER
            if ri == 3:
                c.fill = HEADER_FILL; c.font = HEADER_FONT
                c.alignment = Alignment(wrap_text=True)
            elif ci == 9 and isinstance(val,int) and val > 0:  # columna 118
                c.fill = CRITICAL_FILL; c.font = Font(color="FFFFFF", bold=True)
    ws3['A10'] = "Lectura: a z=12-15, 118 (columna 9, NUNCA debe activarse) dispara MÁS eventos que 217 mismo (columna 11, DEBE activarse). A z=20-25, 217 cae a 0 (se pierde). No hay zona limpia."
    ws3['A10'].font = Font(italic=True, size=9)
    for i, w in enumerate([6,7,7,7,7,7,7,7,26,7,20], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    out_path = os.path.join(OUT_DIR, 'tabla_comparativa_3a_ronda.xlsx')
    wb.save(out_path)
    print(f"xlsx guardado en: {out_path}")


if __name__ == '__main__':
    build()
