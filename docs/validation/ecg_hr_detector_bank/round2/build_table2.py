import os, json, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_bank2 import RECORDS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

HEADER_FILL = PatternFill(start_color="1d4ed8", end_color="1d4ed8", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GOOD_FILL = PatternFill(start_color="c6f6d5", end_color="c6f6d5", fill_type="solid")
BAD_FILL = PatternFill(start_color="fed7d7", end_color="fed7d7", fill_type="solid")
MID_FILL = PatternFill(start_color="fefcbf", end_color="fefcbf", fill_type="solid")
CRITICAL_FILL = PatternFill(start_color="dc2626", end_color="dc2626", fill_type="solid")
THIN = Side(style="thin", color="cccccc")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("record", "Registro"), ("label", "Descripción"), ("prominence_factor", "prominence_factor"),
    ("n_gt_beats", "Latidos GT"), ("hr_ground_truth", "HR GT (bpm)"),
    ("sens_old", "Sens. VIEJO"), ("ppv_old", "VPP VIEJO"), ("hr_error_old", "Err.HR VIEJO (NO usar solo)"),
    ("sens_new", "Sens. REPARADO"), ("ppv_new", "VPP REPARADO"), ("hr_error_new", "Err.HR REPARADO (NO usar solo)"),
    ("n_competitions", "Competencias bifásicas"), ("confidence", "Confianza compuerta"), ("gate_decision", "Decisión compuerta"),
]


def quality_fill(value):
    if value is None:
        return None
    if value >= 0.95:
        return GOOD_FILL
    if value >= 0.8:
        return MID_FILL
    return BAD_FILL


def build():
    rows = json.load(open(os.path.join(OUT_DIR, 'raw_results2.json')))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparativa VIEJO vs REPARADO"

    ws.merge_cells('A1:N1')
    ws['A1'] = "Banco de validación 2ª ronda — Detector VIEJO vs REPARADO (bidireccional + refractario) + compuerta de calidad"
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells('A2:N2')
    ws['A2'] = ("Ground truth = anotaciones .atr de MIT-BIH · Tolerancia ±50ms · Ventana 180s · "
                "⚠ el error de HR está marcado 'NO usar solo' -- úsense sensibilidad/VPP como criterio")
    ws['A2'].font = Font(italic=True, size=9, color="c0392b")

    header_row = 4
    for i, (key, label) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BORDER

    r = header_row + 1
    for row in rows:
        for i, (key, label) in enumerate(COLUMNS, start=1):
            val = row[key]
            c = ws.cell(row=r, column=i, value=val)
            c.border = BORDER
            if key in ("sens_old", "ppv_old", "sens_new", "ppv_new"):
                fill = quality_fill(val)
                if fill:
                    c.fill = fill
            if key == "gate_decision":
                c.fill = GOOD_FILL if val == 'ACEPTA' else BAD_FILL
            # marca critica: compuerta ACEPTA pero sens_new muy bajo (falso-aceptado)
            if key == "gate_decision" and val == 'ACEPTA' and row.get('sens_new') is not None and row['sens_new'] < 0.5:
                c.fill = CRITICAL_FILL
                c.font = Font(color="FFFFFF", bold=True)
        r += 1

    widths = [10, 46, 16, 10, 12, 11, 10, 16, 14, 11, 18, 13, 13, 13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # Hoja 2: alerta critica aislada
    ws2 = wb.create_sheet("ALERTA CRÍTICA — falso-aceptado")
    ws2['A1'] = "Falso-aceptado de la compuerta de calidad (confianza alta, sensibilidad real muy baja)"
    ws2['A1'].font = Font(bold=True, size=13, color="c0392b")
    headers2 = ["Registro", "prominence_factor", "Sensibilidad real", "VPP real", "Confianza (compuerta)", "Decisión", "Mecanismo"]
    for i, h in enumerate(headers2, start=1):
        c = ws2.cell(row=3, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True)
    r2 = 4
    for row in rows:
        if row['gate_decision'] == 'ACEPTA' and row.get('sens_new') is not None and row['sens_new'] < 0.5:
            vals = [row['record'], row['prominence_factor'], row['sens_new'], row['ppv_new'], row['confidence'], row['gate_decision'],
                    "Ritmo con marcapasos: la supresión refractoria elige la deflexión negativa (mayor amplitud absoluta) del QRS ancho en vez de la positiva que coincide con la anotación -- el desfase resultante (~60-70ms) cae fuera de la tolerancia ±50ms. La regularidad/consistencia de amplitud de las espigas de marcapasos engaña a la métrica de confianza (que no usa ground truth)."]
            for i, v in enumerate(vals, start=1):
                c = ws2.cell(row=r2, column=i, value=v)
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i == 6:
                    c.fill = CRITICAL_FILL
                    c.font = Font(color="FFFFFF", bold=True)
            r2 += 1
    widths2 = [10, 16, 14, 10, 16, 12, 70]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    out_path = os.path.join(OUT_DIR, 'tabla_comparativa_2a_ronda.xlsx')
    wb.save(out_path)
    print(f"xlsx guardado en: {out_path}")


if __name__ == '__main__':
    build()
