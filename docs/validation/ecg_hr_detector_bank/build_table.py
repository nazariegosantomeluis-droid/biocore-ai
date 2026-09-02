import os, json, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_bank import RECORDS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.dirname(os.path.abspath(__file__))

HEADER_FILL = PatternFill(start_color="1d4ed8", end_color="1d4ed8", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
GOOD_FILL = PatternFill(start_color="c6f6d5", end_color="c6f6d5", fill_type="solid")
BAD_FILL = PatternFill(start_color="fed7d7", end_color="fed7d7", fill_type="solid")
MID_FILL = PatternFill(start_color="fefcbf", end_color="fefcbf", fill_type="solid")
THIN = Side(style="thin", color="cccccc")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLUMNS = [
    ("record", "Registro"), ("label", "Descripción"), ("prominence_factor", "prominence_factor"),
    ("n_gt_beats", "Latidos GT (cardiólogo)"), ("n_detected", "Picos detectados"),
    ("hr_ground_truth", "HR ground truth (bpm)"), ("hr_detected", "HR detectado (bpm)"),
    ("hr_error_bpm", "Error HR (bpm)"), ("interval_std_ms", "Desv. std intervalo (ms)"),
    ("tp", "Aciertos (TP)"), ("fp", "Falsos+ (FP)"), ("fn", "Falsos- (FN)"),
    ("sensitivity", "Sensibilidad"), ("ppv", "VPP (precisión)"),
]


def cell_fill_for(col_key, value):
    if value is None:
        return None
    if col_key == "sensitivity" or col_key == "ppv":
        if value >= 0.95:
            return GOOD_FILL
        if value >= 0.8:
            return MID_FILL
        return BAD_FILL
    if col_key == "hr_error_bpm":
        if abs(value) <= 2:
            return GOOD_FILL
        if abs(value) <= 10:
            return MID_FILL
        return BAD_FILL
    return None


def build():
    rows = json.load(open(os.path.join(OUT_DIR, 'raw_results.json')))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Comparativa detector HR"

    ws.merge_cells('A1:N1')
    ws['A1'] = "Banco de validación — Detector de HR de ECG (clinical.ecg_analyzer.ECGAnalyzer.detect_r_peaks)"
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A2:N2')
    ws['A2'] = "Ground truth = anotaciones de cardiólogos (archivos .atr de MIT-BIH, PhysioNet) · Tolerancia de emparejamiento: ±50ms · Ventana de análisis: 180s por registro"
    ws['A2'].font = Font(italic=True, size=9, color="555555")

    header_row = 4
    for i, (key, label) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=header_row, column=i, value=label)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BORDER

    r = header_row + 1
    for row in rows:
        for i, (key, label) in enumerate(COLUMNS, start=1):
            val = row[key]
            c = ws.cell(row=r, column=i, value=val)
            c.border = BORDER
            fill = cell_fill_for(key, val)
            if fill:
                c.fill = fill
        r += 1

    widths = [9, 46, 16, 15, 15, 16, 15, 12, 16, 10, 10, 10, 12, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # Segunda hoja: resumen por registro al pf candidato del diagnóstico (1.0) vs default (0.3)
    ws2 = wb.create_sheet("Resumen 0.3 vs 1.0")
    ws2.merge_cells('A1:H1')
    ws2['A1'] = "Contraste: prominence_factor por defecto (0.3) vs candidato del diagnóstico (1.0)"
    ws2['A1'].font = Font(bold=True, size=13)
    headers2 = ["Registro", "Descripción", "HR gt", "HR@0.3", "Err@0.3", "HR@1.0", "Err@1.0", "Veredicto visual (ver PDF de tiras)"]
    for i, h in enumerate(headers2, start=1):
        c = ws2.cell(row=3, column=i, value=h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True)

    by_rec = {}
    for row in rows:
        by_rec.setdefault(row['record'], {})[row['prominence_factor']] = row

    veredicto = {
        '100': 'Limpio, ambos parámetros funcionan; 1.0 es exacto',
        '103': 'Limpio, 1.0 casi exacto (5 FP), 1.5 exacto',
        '117': 'QRS predominantemente negativo en esta derivación -- find_peaks() (solo picos positivos) NUNCA ve el latido real, en NINGÚN prominence_factor probado (VPP tope ~0.53)',
        '108': 'Mismo problema que 117 (QRS negativo) + bajo voltaje -- sensibilidad catastrófica (0.15-0.18) en todo el barrido',
        '228': 'Arritmia real: 1.0 da el menor error de HR pero sensibilidad/VPP imperfectos (0.85/0.87); umbrales altos pierden latidos ectópicos',
        '207': 'Caso extremo (fibrilación/flutter ventricular + ruido): sensibilidad nunca supera ~0.36 en todo el barrido -- ningún prominence_factor lo resuelve',
        '233': 'Ruidoso pero HR promedio engañosamente preciso (error <3bpm en casi todo el barrido) pese a sensibilidad/VPP imperfectos -- los errores se cancelan en el promedio',
        '118': 'BRD: morfología distinta no afecta el detector -- 1.0/1.5 dan resultados casi perfectos',
    }

    r2 = 4
    for rec_id, label in RECORDS.items():
        d = by_rec.get(rec_id, {})
        row03 = d.get(0.3, {})
        row10 = d.get(1.0, {})
        vals = [
            rec_id, label,
            row03.get('hr_ground_truth'),
            row03.get('hr_detected'), row03.get('hr_error_bpm'),
            row10.get('hr_detected'), row10.get('hr_error_bpm'),
            veredicto.get(rec_id, ''),
        ]
        for i, v in enumerate(vals, start=1):
            c = ws2.cell(row=r2, column=i, value=v)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r2 += 1

    widths2 = [10, 34, 8, 8, 8, 8, 8, 60]
    for i, w in enumerate(widths2, start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    out_path = os.path.join(OUT_DIR, 'tabla_comparativa_detector_hr.xlsx')
    wb.save(out_path)
    print(f"xlsx guardado en: {out_path}")


if __name__ == '__main__':
    build()
